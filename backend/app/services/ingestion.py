"""
Service d'ingestion : parsing des fichiers Excel multi-feuilles.

Étapes :
1. detecter_feuilles()      → liste les feuilles + nb lignes estimé
2. _detecter_header_row()   → détecte la ligne d'entête (auto ou manuel)
3. _analyser_entetes()      → détecte les colonnes (regex)
4. _extraire_lignes()       → extrait et normalise chaque ligne
5. ingerer_feuille()        → insère dans ligne_ingeree
6. ingerer_fichier()        → orchestre sur toutes les feuilles

Modes : complet / echantillon / plage
"""
from __future__ import annotations

import hashlib
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Feuille, Fichier, IngestionRun, LigneIngeree
from app.services.coordonnees import parser_coordonnees

logger = logging.getLogger("dhr.ingestion")


# =====================================================================
# Patterns regex pour la détection des colonnes
# =====================================================================
COLUMN_PATTERNS: dict[str, list[re.Pattern]] = {
    # --- Contexte administratif ---
    "region_brute": [
        re.compile(r"^r[ée]gion$", re.I),
        re.compile(r"q1\.2\.\s*r[ée]gion", re.I),
    ],
    "departement_brute": [
        re.compile(r"^d[ée]partement$", re.I),
        re.compile(r"q1\.3\.\s*d[ée]partement", re.I),
    ],
    "sous_prefecture_brute": [
        re.compile(r"^sous[_\s-]?pr[ée]fecture", re.I),
        re.compile(r"commune\s*/\s*sous[_\s-]?pr[ée]fecture", re.I),
        re.compile(r"q1\.4\.", re.I),
    ],
    "type_localite_brut": [
        re.compile(r"^type[_\s]localit[ée]", re.I),
        re.compile(r"^localit[ée][_\s]ou[_\s]campement$", re.I),  # ← ajouter
    ],
    "localite_brute": [
        re.compile(r"^localit[ée]$", re.I),   # ← seulement "Localité" seul
        re.compile(r"^village$", re.I),
        re.compile(r"^nom[_\s]village$", re.I),
        re.compile(r"q1\.5\.", re.I),
        re.compile(r"q1\.6\.", re.I),
    ],
    # --- Coordonnées ---
    "latitude_decimal": [
        re.compile(r"^latitude$", re.I),
        re.compile(r"^latitude\s*ouvrage$", re.I),
        re.compile(r"^lat$", re.I),
        re.compile(r"^y$", re.I),
    ],
    "longitude_decimal": [
        re.compile(r"^longitude$", re.I),
        re.compile(r"^longitude\s*ouvrage$", re.I),
        re.compile(r"^lon(g)?$", re.I),
        re.compile(r"^x$", re.I),
    ],
    # --- Ouvrage ---
    "type_ouvrage_brut": [
        re.compile(r"^type[_\s]ouvrage$", re.I),
        re.compile(r"^type_?\s*ouvrage$", re.I),
        re.compile(r"q3\.1\.\d+[a-z]?\.\s*type\s*d.?ouvrage", re.I),
    ],
    "etat_ouvrage_brut": [
        re.compile(r"^etat(_ouvrage)?$", re.I),
        re.compile(r"^etat\s*ouvrage$", re.I),
        re.compile(r"^état(_ouvrage)?$", re.I),
        re.compile(r"^etat[_\s]forage$", re.I),
        re.compile(r"q3\.1\.\d+[a-z]?\.\s*[ée]tat\s*du\s*forage", re.I),
    ],
    "annee_projet": [
        re.compile(r"^ann[ée]e", re.I),
        re.compile(r"date[_\s]fin", re.I),
        re.compile(r"date[_\s]r[ée]alisation", re.I),
    ],
    "bailleur_brut": [
        re.compile(r"^bailleur", re.I),
        re.compile(r"^projet$", re.I),
        re.compile(r"^financeur", re.I),
    ],
    "type_financement_brut": [
        re.compile(r"^type[_\s]financement", re.I),
    ],
    "a_cgpe": [
        re.compile(r"^a[_\s]cgpe$", re.I),
        re.compile(r"existence[_\s]cgpe", re.I),
        re.compile(r"^nb[_\s]comite$", re.I),
        re.compile(r"^organe\s*gestion$", re.I),
    ],
}


# =====================================================================
# Helpers
# =====================================================================
def _normaliser(texte: str | None) -> str:
    """Normalise un texte (minuscules, sans accents, espaces compactés)."""
    if texte is None:
        return ""
    texte = str(texte).strip().lower()
    texte = "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )
    texte = re.sub(r"[\s\-_]+", " ", texte).strip()
    return texte


def _hash_fichier(chemin: str | Path) -> str:
    h = hashlib.sha256()
    with open(chemin, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _to_int(val: Any) -> int | None:
    if val is None or val == "":
        return None
    try:
        return int(float(str(val).replace(",", ".").strip()))
    except (ValueError, TypeError):
        return None


def _to_bool(val: Any) -> bool | None:
    if val is None or val == "":
        return None
    s = _normaliser(str(val))
    if s in ("oui", "yes", "true", "1", "vrai"):
        return True
    if s in ("non", "no", "false", "0", "faux"):
        return False
    return None

def _to_float(s: str) -> float | None:
    """Convertit une chaîne en float, gère la virgule. Retourne None si invalide."""
    try:
        return float(s.replace(",", ".").strip())
    except (ValueError, AttributeError):
        return None

def _serialiser_valeur(val: Any) -> Any:
    from datetime import date, datetime, time
    if val is None:
        return None
    if isinstance(val, (datetime, date, time)):
        return val.isoformat()
    if isinstance(val, (int, float, str, bool)):
        return val
    # Fallback : str
    return str(val)

# =====================================================================
# 1. Détection des feuilles
# =====================================================================
@dataclass
class FeuilleInfo:
    index: int
    nom: str
    nb_lignes: int
    nb_colonnes: int
    header_row: int = 1
    entetes: list[str] = field(default_factory=list)


def _detecter_header_row(ws, max_scan: int = 30) -> int:
    """
    Détecte la ligne d'entête :
    - Cherche la première ligne avec >= 5 cellules non vides
    - Ignore les lignes de séparateurs (tirets/underscores/égal)
    - Ignore les lignes de titre (1-2 cellules non vides)
    - Score basé sur le nombre de colonnes uniques
    """
    best_row = 1
    best_score = 0

    for row_idx in range(1, min(max_scan + 1, (ws.max_row or 1) + 1)):
        row = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
        non_vides = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]

        # Ignore les lignes vides ou avec trop peu de cellules
        if len(non_vides) < 5:
            continue

        # Ignore les séparateurs (tirets/underscores/égal)
        if all(re.fullmatch(r"[-_=*]+", v) for v in non_vides):
            continue

        # Score = nombre de colonnes uniques
        uniques = len(set(non_vides))
        score = uniques

        if score > best_score:
            best_score = score
            best_row = row_idx

        # Si on a une ligne avec >= 10 colonnes uniques, on s'arrête
        if uniques >= 10:
            return row_idx

    return best_row

def detecter_feuilles(chemin: str | Path, header_row_force: int | None = None) -> list[FeuilleInfo]:
    """Ouvre un fichier Excel et liste ses feuilles."""
    try:
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Impossible d'ouvrir le fichier Excel : {e}",
        )
    try:
        feuilles: list[FeuilleInfo] = []
        for idx, nom in enumerate(wb.sheetnames, start=1):
            ws = wb[nom]
            nb_lignes_total = ws.max_row or 0
            nb_colonnes = ws.max_column or 0

            header_row = header_row_force or _detecter_header_row(ws)

            entetes: list[str] = []
            if nb_lignes_total >= header_row:
                for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=False)):
                    entetes.append(str(cell.value) if cell.value is not None else "")

            feuilles.append(FeuilleInfo(
                index=idx,
                nom=nom,
                nb_lignes=max(0, nb_lignes_total - header_row),
                nb_colonnes=nb_colonnes,
                header_row=header_row,
                entetes=entetes,
            ))
        return feuilles
    finally:
        wb.close()


# =====================================================================
# 2. Mapping des colonnes
# =====================================================================
@dataclass
class MappingFeuille:
    index_feuille: int
    header_row: int
    colonnes_mappees: dict[int, str] = field(default_factory=dict)
    colonnes_non_mappees: list[str] = field(default_factory=list)
    entetes: list[str] = field(default_factory=list)


def _analyser_entetes(entetes: list[str], index_feuille: int, header_row: int) -> MappingFeuille:
    mapping = MappingFeuille(index_feuille=index_feuille, header_row=header_row, entetes=entetes)
    champs_deja_mappes: set[str] = set()

    for col_idx, entete in enumerate(entetes):
        if not entete or not entete.strip():
            continue
        entete_norm = _normaliser(entete)
        matched_champ = None
        for champ, patterns in COLUMN_PATTERNS.items():
            if champ in champs_deja_mappes:
                continue
            for pat in patterns:
                if pat.search(entete_norm):
                    matched_champ = champ
                    break
            if matched_champ:
                break

        if matched_champ:
            mapping.colonnes_mappees[col_idx] = matched_champ
            champs_deja_mappes.add(matched_champ)
        else:
            mapping.colonnes_non_mappees.append(entete)

    return mapping


# =====================================================================
# 3. Extraction des lignes
# =====================================================================
def _extraire_lignes(
    ws,
    mapping: MappingFeuille,
    mode: str,
    limite_lignes: int | None,
    ligne_debut: int | None,
    ligne_fin: int | None,
) -> list[dict[str, Any]]:
    """Parcourt les lignes et normalise. Itération unique (optimisé read_only)."""
    lignes: list[dict[str, Any]] = []

    start = mapping.header_row + 1
    end = None

    if mode == "complet":
        end = ws.max_row
    elif mode == "echantillon":
        end = min(start + (limite_lignes or 100) - 1, ws.max_row)
    elif mode == "plage":
        start = max(mapping.header_row + 1, (ligne_debut or start))
        end = min(ligne_fin or ws.max_row, ws.max_row)
    else:
        raise HTTPException(status_code=400, detail=f"Mode invalide : {mode}")

    # Itération UNIQUE sur toutes les lignes (optimisé read_only)
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True),
        start=1,
    ):
        if row_idx < start:
            continue
        if end is not None and row_idx > end:
            break

        # Ignore les lignes vides
        if all(v is None or str(v).strip() == "" for v in row):
            continue

        record: dict[str, Any] = {"num_ligne_source": row_idx}

        for col_idx, champ in mapping.colonnes_mappees.items():
            if col_idx < len(row):
               val = row[col_idx]
               if val is not None and str(val).strip() != "":
                   if champ not in record or record[champ] is None or str(record[champ]).strip() == "":
                      record[champ] = val

        # Conversion coordonnées
        lat_brut = record.get("latitude_decimal")
        lon_brut = record.get("longitude_decimal")
        lat_dec, lon_dec = parser_coordonnees(lat_brut, lon_brut)
        record["latitude_decimal"] = lat_dec
        record["longitude_decimal"] = lon_dec

        # Données brutes
        brut: dict[str, Any] = {}
        for col_idx, val in enumerate(row):
            if col_idx < len(mapping.entetes) and val is not None and mapping.entetes[col_idx]:
                brut[mapping.entetes[col_idx]] = _serialiser_valeur(val)
        record["donnees_brutes"] = brut

        lignes.append(record)

    return lignes


# =====================================================================
# 4. Ingestion d'une feuille
# =====================================================================
def ingerer_feuille(
    db: Session,
    feuille: Feuille,
    run: IngestionRun,
    mode: str,
    limite_lignes: int | None,
    ligne_debut: int | None,
    ligne_fin: int | None,
    chemin_fichier: str,
    header_row_force: int | None = None,
) -> int:
    """Ingère une feuille dans ligne_ingeree."""
    if not Path(chemin_fichier).exists():
        raise HTTPException(
            status_code=404,
            detail=f"Fichier introuvable : {chemin_fichier}",
        )

    wb = openpyxl.load_workbook(chemin_fichier, read_only=True, data_only=True)
    try:
        if feuille.nom_feuille not in wb.sheetnames:
            raise HTTPException(404, f"Feuille '{feuille.nom_feuille}' introuvable")
        ws = wb[feuille.nom_feuille]

        header_row = header_row_force or _detecter_header_row(ws)

        entetes: list[str] = []
        for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=False)):
            entetes.append(str(cell.value) if cell.value is not None else "")

        mapping = _analyser_entetes(entetes, feuille.index_feuille, header_row)

        lignes = _extraire_lignes(ws, mapping, mode, limite_lignes, ligne_debut, ligne_fin)

        # Fallback type_ouvrage depuis le nom de la feuille
        nom_feuille_upper = feuille.nom_feuille.upper()
        type_feuille = None
        for code in ("PMH", "HVA", "PEA", "HV"):
            if code in nom_feuille_upper:
                type_feuille = code
                break

        nb = 0
        for ligne in lignes:
            type_ouvrage = ligne.get("type_ouvrage_brut") or type_feuille
            record = LigneIngeree(
                id_feuille=feuille.id_feuille,
                id_run=run.id_run,
                num_ligne_source=ligne["num_ligne_source"],
                donnees_brutes=ligne["donnees_brutes"],
                district_brut=ligne.get("district_brut"),
                region_brute=ligne.get("region_brute"),
                departement_brut=ligne.get("departement_brute"),
                sous_prefecture_brute=ligne.get("sous_prefecture_brute"),
                localite_brute=ligne.get("localite_brute"),
                type_localite_brut=ligne.get("type_localite_brut"),
                latitude_decimal=ligne.get("latitude_decimal"),
                longitude_decimal=ligne.get("longitude_decimal"),
                type_ouvrage_brut=type_ouvrage,
                etat_ouvrage_brut=ligne.get("etat_ouvrage_brut"),
                annee_projet=_to_int(ligne.get("annee_projet")),
                bailleur_brut=ligne.get("bailleur_brut"),
                type_financement_brut=ligne.get("type_financement_brut"),
                a_cgpe=_to_bool(ligne.get("a_cgpe")),
                statut_qualite="ok",
            )
            db.add(record)
            nb += 1

        db.flush()
        feuille.nb_lignes = len(lignes)
        feuille.nb_lignes_valides = nb
        feuille.statut_ingestion = "ingere"
        db.flush()
        return nb
    finally:
        wb.close()


# =====================================================================
# 5. Ingestion d'un fichier
# =====================================================================
def ingerer_fichier(
    db: Session,
    fichier: Fichier,
    run: IngestionRun,
    mode: str = "complet",
    limite_lignes: int | None = None,
    ligne_debut: int | None = None,
    ligne_fin: int | None = None,
    header_row_force: int | None = None,
) -> dict[str, Any]:
    """Orchestre l'ingestion d'un fichier."""
    from datetime import datetime, timezone

    feuilles = db.scalars(select(Feuille).where(Feuille.id_fichier == fichier.id_fichier)).all()
    if not feuilles:
        raise HTTPException(400, f"Aucune feuille pour le fichier {fichier.id_fichier}")

    total = 0
    resultats = []

    for feuille in feuilles:
        # Capture les attributs AVANT le try (évite PendingRollbackError)
        feuille_id = feuille.id_feuille
        feuille_nom = feuille.nom_feuille

        try:
            nb = ingerer_feuille(
                db, feuille, run, mode, limite_lignes, ligne_debut, ligne_fin,
                fichier.chemin_stockage, header_row_force,
            )
            total += nb
            resultats.append({
                "id_feuille": feuille_id,
                "nom_feuille": feuille_nom,
                "statut": "ok",
                "nb_lignes": nb,
            })
        except Exception as e:
            import traceback
            # Rollback AVANT d'accéder aux attributs ORM
            db.rollback()
            logger.exception("Erreur feuille %s : %s", feuille_nom, e,) #  traceback.format_exc()
            # Recharge la feuille (après rollback)
            feuille = db.get(Feuille, feuille_id)
            if feuille is not None:
                feuille.statut_ingestion = "erreur"
                feuille.message_erreur = str(e)
                db.flush()
            resultats.append({
                "id_feuille": feuille_id,
                "nom_feuille": feuille_nom,
                "statut": "erreur",
                "message": str(e),
            })

    run.nb_lignes_ingerees = total
    run.statut = "termine" if total > 0 else "erreur"
    run.finished_at = datetime.now(timezone.utc)
    fichier.nb_lignes_total = total
    fichier.statut = "ingere"
    fichier.date_ingestion = datetime.now(timezone.utc)
    db.flush()

    return {
        "id_fichier": fichier.id_fichier,
        "id_run": run.id_run,
        "total_lignes": total,
        "feuilles": resultats,
    }