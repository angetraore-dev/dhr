"""Inventaire des fichiers Excel."""
import sys
from pathlib import Path
import openpyxl


def inventaire(chemin: str) -> None:
    path = Path(chemin)
    if not path.exists():
        print(f"[X] Fichier introuvable : {chemin}")
        return
    print(f"\n{'='*70}")
    print(f"Fichier : {path.name} ({path.stat().st_size / 1024:.1f} Ko)")
    print(f"{'='*70}")
    try:
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    except Exception as e:
        print(f"[X] Erreur ouverture : {e}")
        return
    try:
        for nom in wb.sheetnames:
            ws = wb[nom]
            nb_lignes = ws.max_row or 0
            nb_cols = ws.max_column or 0
            print(f"\n  Feuille '{nom}' : {nb_lignes} lignes, {nb_cols} colonnes")
            if nb_lignes < 1:
                continue
            entetes = []
            for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=False)):
                entetes.append(str(cell.value) if cell.value is not None else "")
            print(f"     Entetes ({len(entetes)}) :")
            for i, h in enumerate(entetes):
                print(f"       [{i}] {h!r}")
            print(f"     Exemples de donnees :")
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, max_row=min(4, nb_lignes), values_only=True),
                start=1,
            ):
                row_str = []
                for v in row:
                    s = str(v) if v is not None else ""
                    if len(s) > 40:
                        s = s[:37] + "..."
                    row_str.append(s)
                print(f"       Ligne {row_idx} : {row_str}")
    finally:
        wb.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage : python inventaire_fichiers.py <fichier.xlsx> [fichier2.xlsx ...]")
        sys.exit(1)
    for chemin in sys.argv[1:]:
        inventaire(chemin)
